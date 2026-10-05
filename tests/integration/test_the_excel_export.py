"""The workbook an operator takes away: every order, every customer."""

from __future__ import annotations

import io
import uuid
from collections.abc import Iterator
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from geekvpn.infrastructure.analytics.xlsx_export import build_workbook
from geekvpn.infrastructure.config.settings import get_settings
from geekvpn.infrastructure.persistence.base import Base
from geekvpn.infrastructure.persistence.models.identity import UserModel
from geekvpn.infrastructure.persistence.models.payments import WalletEntryModel
from geekvpn.infrastructure.persistence.models.provisioning import OrderModel

pytestmark = pytest.mark.integration

NOW = datetime(2026, 10, 5, 12, 0, tzinfo=UTC)


@pytest.fixture
def session() -> Iterator[Session]:
    engine = create_engine(get_settings().postgres.dsn(driver="postgresql+psycopg"))
    try:
        with engine.connect() as probe:
            probe.execute(text("SELECT 1"))
    except OperationalError as exc:
        pytest.skip(f"no Postgres available: {exc.__class__.__name__}")
    with engine.begin() as connection:
        connection.execute(text("DROP SCHEMA public CASCADE"))
        connection.execute(text("CREATE SCHEMA public"))
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        db.add(
            UserModel(
                id=uuid.uuid4(),
                telegram_id=1001,
                username="ali",
                first_name="علی",
                status="active",
                referral_code="r1",
                language="fa",
            )
        )
        for number, state in (("O-1", "active"), ("O-2", "cancelled")):
            db.add(
                OrderModel(
                    id=uuid.uuid4().hex,
                    number=number,
                    user_id=1001,
                    state=state,
                    plan_id=uuid.uuid4(),
                    product_id=uuid.uuid4(),
                    plan_name_fa="ماهانه",
                    duration_days=30,
                    list_price=200_000,
                    total=200_000,
                    placed_at=NOW - timedelta(days=1),
                )
            )
        db.add(
            WalletEntryModel(
                id="w1",
                user_id=1001,
                kind="adjustment",
                amount=70_000,
                balance_after=70_000,
                occurred_at=NOW,
                description_fa="هدیه",
            )
        )
        db.commit()
        yield db
    engine.dispose()


def test_both_sheets_carry_the_rows_an_operator_expects(session: Session) -> None:
    book = load_workbook(io.BytesIO(build_workbook(session, now=NOW)))

    assert book.sheetnames == ["سفارش‌ها", "کاربران"]
    orders = list(book["سفارش‌ها"].iter_rows(min_row=2, values_only=True))
    assert sorted(row[0] for row in orders) == ["O-1", "O-2"]

    [customer] = list(book["کاربران"].iter_rows(min_row=2, values_only=True))
    # Telegram id, wallet balance, successful purchases, their total.
    assert customer[0] == 1001
    assert customer[1] == "@ali"
    assert (customer[4], customer[5], customer[6]) == (70_000, 1, 200_000)
